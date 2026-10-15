#!/usr/bin/env python3
"""
Tổng hợp kết quả thí nghiệm vào một file xlsx: Reports/ketqua_thi_nghiem.xlsx

  python utils/export_results_xlsx.py

Quét:
  Reports/Meeting_*/checkpoints/<exp>/<run>/config.yaml      → data dùng để train
  Reports/Meeting_*/results/<run>/eval_data_<x>/eval_*.txt   → điểm (JSON sau dòng "[split] N ảnh")

Hai tab:
  "Kết quả"  ngắn gọn: mỗi dòng là (run, data chấm), các cột tự điền + 3 cột nhập tay (Mã, Đổi gì, Ghi chú).
  "Chi tiết" đủ 7 chỉ số của từng split, tự sinh hoàn toàn (dùng khi viết luận văn).

Chạy lại bất cứ lúc nào: các cột nhập tay của tab "Kết quả" được giữ nguyên (khớp theo Run + Data chấm);
dòng của run mới tự xuất hiện với phần nhập tay để trống. Chỉ ĐỌC các thư mục kết quả, không sửa gì ngoài file xlsx.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = Path("Reports/ketqua_thi_nghiem.xlsx")
SHEET_MAIN, SHEET_DETAIL = "Kết quả", "Chi tiết"
MAIN_COLS = ["Mã", "Run", "Đổi gì", "Data train", "Data chấm", "e2e_f1 test (ghép)", "e2e_f1 test_real (ảnh thật)",
             "Sót con bệnh (recall test_real)", "Ghi chú"]
MANUAL = ("Mã", "Đổi gì", "Ghi chú")
METRICS = ["inst_precision", "inst_recall", "inst_f1", "cls_acc_matched", "e2e_f1", "grasserie_recall", "body_dice"]
SPLIT_ORDER = ["valid", "test", "valid_real", "test_real"]
EVAL_RE = re.compile(r"eval_(?:new_|old_)?(.+)\.txt$")

# Tên dễ hiểu của thư mục data. Tên lạ thì ghi nguyên.
DATA_NAMES = {"cu": "legacy/sam3_aug20k (cũ)", "v3": "sam3_aug20k_v3"}
# Run có config.yaml không ghi đúng data đã train (ghi tay)
TRAIN_DATA_OVERRIDE = {"2026-10-03_23-12-56": "sam3_aug20k_v1 (bị rò rỉ)"}

# Điền sẵn phần nhập tay cho các run đã có. Khóa: (run, data chấm). Có sẵn trong file xlsx thì file xlsx thắng.
INITIAL_MANUAL = {
    ("2026-10-14_00-01-58", DATA_NAMES["cu"]): (
        "B0", "Baseline VM-UNet, train từ đầu (chưa có pretrained)", "Chấm trên data cũ (mask chưa sửa)"),
    ("2026-10-14_00-01-58", DATA_NAMES["v3"]): (
        "B0", "(như trên)", "Chấm lại trên data mới để so công bằng với B1"),
    ("2026-10-14_23-36-58", DATA_NAMES["v3"]): (
        "B1", "Sửa mask bằng SAM 3 (data v3), model giữ nguyên",
        "Baseline hiện tại. test thấp vì quy ước một con một id (trần 0.79), xem DATASET_STRUCTURE §2.5"),
}
# Run không có file eval_*.txt nhưng cần ghi lại (điểm lấy từ experiments/multitask_vmunet/notes.md)
EXTRA_ROWS = [{
    "Mã": "—", "Run": "2026-10-03_23-12-56", "Đổi gì": "Lần train đầu của model per-larva",
    "Data train": TRAIN_DATA_OVERRIDE["2026-10-03_23-12-56"], "Data chấm": "sam3_aug20k_v1 (bị rò rỉ)",
    "e2e_f1 test (ghép)": 0.847, "e2e_f1 test_real (ảnh thật)": 0.899, "Sót con bệnh (recall test_real)": None,
    "Ghi chú": "Data rò rỉ ~80% valid/test nên không so sánh được (valid e2e_f1 0.872)",
}]


def read_eval(path: Path) -> dict[str, float]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return json.loads(text[text.index("{"):text.rindex("}") + 1])


def train_data_of(run: str, configs: dict[str, dict]) -> str:
    if run in TRAIN_DATA_OVERRIDE:
        return TRAIN_DATA_OVERRIDE[run]
    root = configs.get(run, {}).get("data_root")
    return Path(root).name if root else "?"


def collect() -> tuple[list[dict], list[dict]]:
    configs: dict[str, dict] = {}
    for cfg_p in Path("Reports").glob("Meeting_*/checkpoints/*/*/config.yaml"):
        configs[cfg_p.parent.name] = yaml.safe_load(cfg_p.read_text())

    detail, groups = [], {}
    for ev in sorted(Path("Reports").glob("Meeting_*/results/*/eval_data_*/eval_*.txt")):
        m = EVAL_RE.match(ev.name)
        if not m:
            continue
        run, data_key = ev.parent.parent.name, ev.parent.name.removeprefix("eval_data_")
        data = DATA_NAMES.get(data_key, ev.parent.name)
        scores = read_eval(ev)
        detail.append({"Run": run, "Data chấm": data, "Split": m.group(1), **{k: scores.get(k) for k in METRICS},
                       "n_gt_worms": scores.get("n_gt_worms"), "Nguồn": str(ev),
                       "Quy ước đếm": "một con = một id"})
        groups.setdefault((run, data), {})[m.group(1)] = scores

    main = []
    for (run, data), by_split in groups.items():
        main.append({
            "Mã": "", "Run": run, "Đổi gì": "", "Data train": train_data_of(run, configs), "Data chấm": data,
            "e2e_f1 test (ghép)": by_split.get("test", {}).get("e2e_f1"),
            "e2e_f1 test_real (ảnh thật)": by_split.get("test_real", {}).get("e2e_f1"),
            "Sót con bệnh (recall test_real)": by_split.get("test_real", {}).get("grasserie_recall"),
            "Ghi chú": "",
        })
    main += [dict(r) for r in EXTRA_ROWS if not any(x["Run"] == r["Run"] and x["Data chấm"] == r["Data chấm"] for x in main)]
    main.sort(key=lambda r: (r["Run"], r["Data chấm"]))
    detail.sort(key=lambda r: (r["Run"], r["Data chấm"],
                               SPLIT_ORDER.index(r["Split"]) if r["Split"] in SPLIT_ORDER else 99, r["Split"]))
    return main, detail


def previous_manual() -> dict[tuple[str, str], dict[str, str]]:
    """Đọc phần nhập tay trong file xlsx cũ (nếu có)."""
    if not XLSX.exists():
        return {}
    wb = load_workbook(XLSX)
    if SHEET_MAIN not in wb.sheetnames:
        return {}
    ws = wb[SHEET_MAIN]
    header = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        if rec.get("Run"):
            out[(rec["Run"], rec["Data chấm"])] = {k: (rec.get(k) or "") for k in MANUAL}
    return out


def style(ws, widths: dict[str, int], number_cols: list[str], wrap_cols: list[str]) -> None:
    header = [c.value for c in ws[1]]
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, name in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, i)
            cell.alignment = Alignment(wrap_text=name in wrap_cols, vertical="top")
            if name in number_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "0.000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    main_rows, detail_rows = collect()
    old = previous_manual()
    for r in main_rows:
        key = (r["Run"], r["Data chấm"])
        init = dict(zip(MANUAL, INITIAL_MANUAL.get(key, ("", "", ""))))
        for k in MANUAL:
            r[k] = old.get(key, {}).get(k) or r.get(k) or init[k]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_MAIN
    ws.append(MAIN_COLS)
    for r in main_rows:
        ws.append([r.get(c) for c in MAIN_COLS])
    style(ws, {"Mã": 6, "Run": 22, "Đổi gì": 38, "Data train": 24, "Data chấm": 24, "e2e_f1 test (ghép)": 13,
               "e2e_f1 test_real (ảnh thật)": 15, "Sót con bệnh (recall test_real)": 16, "Ghi chú": 48},
          ["e2e_f1 test (ghép)", "e2e_f1 test_real (ảnh thật)", "Sót con bệnh (recall test_real)"], ["Đổi gì", "Ghi chú"])

    wd = wb.create_sheet(SHEET_DETAIL)
    cols = ["Run", "Data chấm", "Split", *METRICS, "n_gt_worms", "Quy ước đếm", "Nguồn"]
    wd.append(cols)
    for r in detail_rows:
        wd.append([r.get(c) for c in cols])
    style(wd, {"Run": 22, "Data chấm": 24, "Split": 12, "Quy ước đếm": 18, "Nguồn": 70}, METRICS, [])

    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    print(f"Đã ghi {XLSX}: {len(main_rows)} dòng ở tab '{SHEET_MAIN}', {len(detail_rows)} dòng ở tab '{SHEET_DETAIL}'")


if __name__ == "__main__":
    main()
